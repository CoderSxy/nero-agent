"""Run inside a non-root, network-disabled container; never on the server host."""
from pathlib import Path
import subprocess
import sys
import zipfile
from pptx import Presentation

root = Path('/workspace/skill')
root.mkdir()
with zipfile.ZipFile('/bundle/ppt-nero.zip') as archive:
    archive.extractall(root)  # Package is checksummed and validated before upload.
project = Path('/workspace/smoke')
(project / 'svg_output').mkdir(parents=True)
(project / 'notes').mkdir()
(project / 'svg_output' / 'slide_one.svg').write_text('''<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720"><rect width="1280" height="720" fill="#FFFFFF"/><text x="80" y="140" font-family="Noto Sans CJK SC" font-size="40" fill="#111111">PPT 技能导出验证</text></svg>''')
(project / 'notes' / 'total.md').write_text('# slide_one\n\nPPT smoke test.\n')
for name in ['total_md_split.py', 'finalize_svg.py', 'svg_to_pptx.py']:
    subprocess.run([sys.executable, str(root / 'scripts' / name), str(project)], check=True)
outputs = list((project / 'exports').glob('*.pptx'))
assert outputs, 'No PPTX output'
assert len(Presentation(str(outputs[0])).slides) == 1
assert 'PPT' in ''.join(shape.text for shape in Presentation(str(outputs[0])).slides[0].shapes if shape.has_text_frame)
print('PPT_SMOKE_OK: offline SVG -> finalized SVG -> one-slide editable PPTX')
